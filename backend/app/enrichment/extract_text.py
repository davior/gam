"""The extract_text job: a document in, readable text out.

The odd one out of the enrichment set. It calls no provider, spends nothing, and nobody
asks for it because they want to read its output — they want `summarize` and `autotag`
to stop refusing. Those jobs read whatever `enrichment/source.py` hands them, and for a
document it handed them a refusal, because this module did not exist.

Runs on a worker thread, so it owns its session and reports progress through the
callback the queue hands it — which is also the cancellation checkpoint.
"""

from __future__ import annotations

import csv
import io
import logging
from pathlib import Path
from typing import Callable, Iterator, Optional

from sqlmodel import Session, col, delete

from app.ingest.filetypes import TYPE_DOCUMENT, extension_of
from app.models.asset import Asset
from app.models.document import DocumentPage
from app.storage import build_storage

logger = logging.getLogger(__name__)

Progress = Callable[..., None]

# Formats with a unit of their own — a page, a slide, a sheet.
EXT_PDF = ".pdf"
EXT_DOCX = ".docx"
EXT_PPTX = ".pptx"
EXT_XLSX = ".xlsx"
PLAIN_EXTENSIONS = frozenset({".txt", ".md", ".csv"})

EXTRACTABLE_EXTENSIONS = frozenset(
    {EXT_PDF, EXT_DOCX, EXT_PPTX, EXT_XLSX} | PLAIN_EXTENSIONS
)

# How much text goes in one row, for the formats with no unit of their own. A row is the
# chunk that will be embedded individually when full-document search lands, and embedders
# cap somewhere near 8k tokens, so this leaves headroom rather than sitting on the limit.
# Splitting happens on a paragraph boundary where there is one within reach.
MAX_CHUNK_CHARS = 4000

# Past this, stop reading. A document that long is a data dump rather than something
# anyone wants summarised, and the LLM jobs truncate at 60k anyway — the difference this
# bound makes is to the database and to how long the job holds a worker thread.
MAX_TOTAL_CHARS = 2_000_000


class TextExtractionError(Exception):
    """Something the user can be told about."""


def extractable(asset: Asset) -> bool:
    if asset.asset_type != TYPE_DOCUMENT or not asset.storage_key:
        return False
    return extension_of(asset.original_name or "") in EXTRACTABLE_EXTENSIONS


def run(session: Session, asset: Asset, progress: Progress) -> str:
    """Extract one document's text. Returns what to show in the activity feed."""
    if asset.asset_type != TYPE_DOCUMENT or not asset.storage_key:
        raise TextExtractionError("Only documents can have their text extracted")

    extension = extension_of(asset.original_name or "")
    if extension not in EXTRACTABLE_EXTENSIONS:
        # Named rather than generic: "we cannot read .rtf" is actionable — convert it and
        # upload again — where "extraction failed" invites a re-run that cannot work.
        raise TextExtractionError(
            f"Reading text out of {extension or 'this format'} files is not supported. "
            "Convert it to PDF, DOCX or plain text and upload it again."
        )

    progress("Reading the file", 5, "")

    storage = build_storage()
    with storage.materialise(asset.storage_key) as path:
        try:
            chunks = list(_extract(path, extension, progress))
        except TextExtractionError:
            raise
        except Exception as exc:  # noqa: BLE001 - upstream parsers raise their own zoo
            raise TextExtractionError(
                "The file could not be read. It may be corrupt, password-protected or "
                "not really the format its name claims."
            ) from exc

    if not chunks:
        raise TextExtractionError(
            "No text could be read from this document. A scan or a photographed page "
            "holds pictures of words rather than words, which needs OCR."
        )

    progress("Saving", 90, "")
    stored = _store(session, asset, chunks)

    unit = _unit_name(extension, stored)
    return f"{stored} {unit}"


class _Chunk:
    __slots__ = ("page_number", "label", "text")

    def __init__(self, text: str, *, page_number: Optional[int] = None, label: Optional[str] = None):
        self.text = text
        self.page_number = page_number
        self.label = label


def _extract(path: Path, extension: str, progress: Progress) -> Iterator[_Chunk]:
    if extension == EXT_PDF:
        yield from _from_pdf(path, progress)
    elif extension == EXT_DOCX:
        yield from _from_docx(path, progress)
    elif extension == EXT_PPTX:
        yield from _from_pptx(path, progress)
    elif extension == EXT_XLSX:
        yield from _from_xlsx(path, progress)
    else:
        yield from _from_plain(path, progress)


def _from_pdf(path: Path, progress: Progress) -> Iterator[_Chunk]:
    """One row per page, which is the one format where that is simply true.

    pypdfium2 rather than pypdf: it is already a dependency for first-page thumbnails,
    and `ingest/thumbnails.py` records why it was chosen over the AGPL and
    poppler-dependent alternatives. Adding a second PDF library to read what this one
    already has open would reopen a settled argument.
    """
    import pypdfium2

    document = pypdfium2.PdfDocument(str(path))
    try:
        total = len(document)
        budget = MAX_TOTAL_CHARS
        for index in range(total):
            progress("Extracting text", _percent(index, total), f"page {index + 1} of {total}")
            page = document[index]
            textpage = page.get_textpage()
            try:
                # `get_text_bounded()`, not `get_text_range()`: the latter warns that a
                # call with default arguments is redirected here anyway, and a warning
                # per page of a long PDF is a lot of noise for the same bytes.
                text = textpage.get_text_bounded()
            finally:
                textpage.close()
                page.close()
            text = _tidy(text)
            if not text:
                continue
            yield _Chunk(text, page_number=index + 1, label=f"Page {index + 1}")
            budget -= len(text)
            if budget <= 0:
                logger.warning("Stopped reading %s at page %s: too much text", path.name, index + 1)
                return
    finally:
        document.close()


def _from_docx(path: Path, progress: Progress) -> Iterator[_Chunk]:
    """Paragraphs, chunked — a .docx has no pages until something renders it.

    Table cells are read as well as paragraphs: a specification or an invoice can carry
    most of its meaning in a table, and skipping them reads as "extraction produced
    nothing" for a file that is visibly full of words.
    """
    import docx

    progress("Extracting text", 20, "")
    document = docx.Document(str(path))

    blocks = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append("\t".join(cells))

    yield from _chunked(blocks, progress)


def _from_pptx(path: Path, progress: Progress) -> Iterator[_Chunk]:
    """One row per slide, including the speaker notes.

    The notes are where the argument usually lives — a slide says "Q3 Revenue" and the
    notes say why it fell — so leaving them out would summarise the headings alone.
    """
    from pptx import Presentation

    presentation = Presentation(str(path))
    slides = list(presentation.slides)
    total = len(slides)

    for index, slide in enumerate(slides):
        progress("Extracting text", _percent(index, total), f"slide {index + 1} of {total}")
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                parts.append(f"Notes: {notes.text}")
        text = _tidy("\n".join(parts))
        if not text:
            continue
        yield _Chunk(text, page_number=index + 1, label=f"Slide {index + 1}")


def _from_xlsx(path: Path, progress: Progress) -> Iterator[_Chunk]:
    """One row per worksheet, cells joined into lines.

    `read_only=True` streams rather than building the whole workbook in memory, which
    matters for the spreadsheets people actually keep. Formulas are not evaluated —
    `data_only=True` would return the last cached value, or None for a sheet never opened
    in Excel, so the formula text is the more honest thing to read.
    """
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=False)
    try:
        sheets = workbook.worksheets
        total = len(sheets)
        for index, sheet in enumerate(sheets):
            progress("Extracting text", _percent(index, total), f"sheet {index + 1} of {total}")
            lines: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(value) for value in row if value is not None]
                if cells:
                    lines.append("\t".join(cells))
            text = _tidy("\n".join(lines))
            if not text:
                continue
            yield _Chunk(
                text,
                page_number=index + 1,
                label=f"Sheet {sheet.title!r}" if sheet.title else f"Sheet {index + 1}",
            )
    finally:
        workbook.close()


def _from_plain(path: Path, progress: Progress) -> Iterator[_Chunk]:
    """Decode and chunk.

    UTF-8 with `errors="replace"`, and no encoding detection: there is no chardet in this
    tree, and guessing wrong silently is worse than a few replacement characters in a
    summary. A CSV is read through the csv module so quoted cells containing newlines do
    not split a record in half.
    """
    progress("Extracting text", 20, "")
    raw = path.read_bytes()
    if len(raw) > MAX_TOTAL_CHARS:
        raw = raw[:MAX_TOTAL_CHARS]
    body = raw.decode("utf-8", errors="replace")

    if extension_of(path.name) == ".csv":
        reader = csv.reader(io.StringIO(body))
        blocks = ["\t".join(row) for row in reader if any(cell.strip() for cell in row)]
    else:
        blocks = body.split("\n")

    yield from _chunked(blocks, progress)


def _chunked(blocks: list[str], progress: Progress) -> Iterator[_Chunk]:
    """Group blocks into rows of at most MAX_CHUNK_CHARS, never splitting a block.

    `page_number` stays null here, deliberately: these formats have no page a reader
    could turn to, and numbering the chunks would put a number in a column whose whole
    contract is that it means the page the reader sees.
    """
    current: list[str] = []
    size = 0
    part = 0
    total_chars = 0

    def flush() -> Optional[_Chunk]:
        nonlocal current, size, part
        text = _tidy("\n".join(current))
        current = []
        size = 0
        if not text:
            return None
        part += 1
        return _Chunk(text, label=f"Part {part}")

    for block in blocks:
        block = block.rstrip()
        if size and size + len(block) > MAX_CHUNK_CHARS:
            chunk = flush()
            if chunk is not None:
                yield chunk
                total_chars += len(chunk.text)
                if total_chars >= MAX_TOTAL_CHARS:
                    return
                progress("Extracting text", min(80, 20 + part), f"part {part}")
        current.append(block)
        size += len(block) + 1

    chunk = flush()
    if chunk is not None:
        yield chunk


def _store(session: Session, asset: Asset, chunks: list[_Chunk]) -> int:
    """Replace what was there.

    Wholesale, unlike a transcript re-run, which preserves hand-edited segments. Nothing
    edits extracted text by hand — there is no editor for it — so there is nothing to
    protect, and keeping old rows around would double the document on every re-run.
    """
    session.exec(delete(DocumentPage).where(col(DocumentPage.asset_id) == asset.id))

    for index, chunk in enumerate(chunks):
        session.add(
            DocumentPage(
                asset_id=asset.id,
                user_id=asset.user_id,
                idx=index,
                page_number=chunk.page_number,
                label=chunk.label,
                text=chunk.text,
            )
        )

    session.commit()
    return len(chunks)


def _tidy(text: str) -> str:
    """Collapse the whitespace a PDF extractor leaves behind, keeping paragraph breaks."""
    lines = [line.strip() for line in (text or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    out: list[str] = []
    for line in lines:
        if line:
            out.append(line)
        elif out and out[-1]:
            out.append("")
    return "\n".join(out).strip()


def _percent(index: int, total: int) -> int:
    if total <= 0:
        return 50
    return 10 + int(70 * index / total)


def _unit_name(extension: str, count: int) -> str:
    if extension == EXT_PDF:
        return "page" if count == 1 else "pages"
    if extension == EXT_PPTX:
        return "slide" if count == 1 else "slides"
    if extension == EXT_XLSX:
        return "sheet" if count == 1 else "sheets"
    return "section" if count == 1 else "sections"
