"""Attribution the file already carries.

A JPEG's EXIF `Artist`, an MP3's ID3 frames, an MP4's tag block, a PDF's `Author`, a
.docx's core properties — all written by whoever produced the file. Reading them is not
inference, which is why this writes directly while an AI proposal has to go through the
suggestion queue. The asymmetry is the point; see docs/m10-attribution.md.

None of it is new data on the wire: `ingest/probe.py` has always run ffprobe with
`-show_format`, which returns the tag block, and GAM parsed out the duration and threw
the rest away.

Everything here degrades rather than fails, and for the same reason the rest of the
ingest pipeline does: this runs after the row is committed, and a file with unreadable
metadata is still a perfectly good asset.

What is deliberately *not* harvested:

- `retrieved_at` — when *you* fetched something is not a fact the file can know.
- `name` — set from the filename at ingest and the user's to change. A container's
  `title` is frequently boilerplate from an encoder, and silently renaming somebody's
  asset on upload is a worse failure than leaving a field blank.
- `description` / `summary` — not attribution, and owned by the enrichment path.
"""

from __future__ import annotations

import logging
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Mapping, Optional

from app.ingest.filetypes import TYPE_AUDIO, TYPE_DOCUMENT, TYPE_IMAGE, TYPE_VIDEO, extension_of

logger = logging.getLogger(__name__)

# EXIF tag numbers, which Pillow returns as an int-keyed mapping.
_EXIF_ARTIST = 0x013B
_EXIF_COPYRIGHT = 0x8298
_EXIF_DATETIME_ORIGINAL = 0x9003
_EXIF_DATETIME = 0x0132

# IPTC records, as Pillow's IptcImagePlugin keys them: (record, dataset).
_IPTC_BYLINE = (2, 80)
_IPTC_CREDIT = (2, 110)
_IPTC_SOURCE = (2, 115)

# Container tag names, in preference order, per attribution field.
#
# `title` is absent on purpose. In a media container it names *this file*, not a
# containing work, and it is very often an encoder's boilerplate — so it maps to neither
# `name` (see the module docstring) nor `source_title`. `album` and `show` genuinely do
# name a containing work, so they are here. Documents are the other way round and are
# handled separately below: a document's `title` is the work's own title, and there is no
# album to stand in for it.
_TAG_SOURCES: dict[str, tuple[str, ...]] = {
    "creator": ("artist", "author", "album_artist", "composer", "director"),
    "publisher": ("publisher", "label", "network", "studio"),
    "source_title": ("album", "show"),
    "license": ("copyright", "license", "rights"),
    "published_date": ("date", "originaldate", "year", "creation_time"),
    "source_url": ("purl", "url", "wxxx", "comment"),
}


def harvest(
    path: Path,
    asset_type: str,
    original_name: str = "",
    probe_tags: Optional[Mapping[str, str]] = None,
) -> dict[str, Any]:
    """Attribution fields read out of this file. Never raises.

    Returns only fields it actually found — the caller decides what to do with them, and
    `services/assets.py` writes only into fields that are still empty, so a harvest can
    never clobber something a person typed.
    """
    try:
        if asset_type in (TYPE_VIDEO, TYPE_AUDIO):
            return _from_container_tags(probe_tags or {})
        if asset_type == TYPE_IMAGE:
            return _from_image(path)
        if asset_type == TYPE_DOCUMENT:
            return _from_document(path, original_name or path.name)
    except Exception as exc:  # noqa: BLE001 - metadata must never fail an upload
        logger.warning("Could not read embedded metadata from %s: %s", path.name, exc)
    return {}


# ─── video and audio ─────────────────────────────────────────────────────────


def _from_container_tags(tags: Mapping[str, str]) -> dict[str, Any]:
    """Map an ffprobe tag block onto attribution fields.

    Only the keys in `_TAG_SOURCES` are read. That allowlist is deliberate: a container
    tag block is mostly technical noise (`encoder`, `handler_name`, `major_brand`,
    `compatible_brands`) and a mapping that took whatever it recognised would file
    "Lavf60.16.100" as somebody's creator.
    """
    found: dict[str, Any] = {}

    for field_name, candidates in _TAG_SOURCES.items():
        for key in candidates:
            value = (tags.get(key) or "").strip()
            if not value:
                continue

            if field_name == "published_date":
                normalised = normalise_partial_date(value)
                if normalised:
                    found[field_name] = normalised
                    break
                continue

            if field_name == "source_url":
                # `comment` is in the candidate list because downloaders routinely put
                # the source URL there — and also because they put everything else
                # there. Only take it when it is actually a URL.
                if _looks_like_url(value):
                    found[field_name] = value
                    break
                continue

            found[field_name] = value
            break

    return found


# ─── images ──────────────────────────────────────────────────────────────────


def _from_image(path: Path) -> dict[str, Any]:
    from PIL import Image

    found: dict[str, Any] = {}
    with Image.open(path) as image:
        exif = image.getexif()
        if exif:
            artist = _clean(exif.get(_EXIF_ARTIST))
            if artist:
                found["creator"] = artist

            rights = _clean(exif.get(_EXIF_COPYRIGHT))
            if rights:
                found["license"] = rights

            # DateTimeOriginal is when the shutter fired; DateTime is when the file was
            # last written, which an edit updates. Prefer the former.
            #
            # DateTimeOriginal lives in the Exif sub-IFD (0x8769), not IFD0, so a
            # top-level lookup alone finds it in hand-built files and misses it in every
            # real photograph — which is the wrong way round for a test to pass.
            sub_ifd = {}
            try:
                sub_ifd = exif.get_ifd(0x8769) or {}
            except Exception:  # noqa: BLE001 - a malformed sub-IFD is not worth failing over
                sub_ifd = {}

            shot = (
                _clean(sub_ifd.get(_EXIF_DATETIME_ORIGINAL))
                or _clean(exif.get(_EXIF_DATETIME_ORIGINAL))
                or _clean(exif.get(_EXIF_DATETIME))
            )
            taken = normalise_partial_date(shot) if shot else None
            if taken:
                found["published_date"] = taken

        found.update(_from_iptc(image))

    return found


def _from_iptc(image: Any) -> dict[str, Any]:
    """IPTC, which is where a press photo's real credit lives.

    EXIF `Artist` is usually the camera owner; a wire photo carries By-line, Credit and
    Source instead, and those are the fields a picture desk actually fills in. They win
    over EXIF for that reason.
    """
    from PIL import IptcImagePlugin

    try:
        info = IptcImagePlugin.getiptcinfo(image)
    except Exception:  # noqa: BLE001 - malformed IPTC is common and not our problem
        return {}
    if not info:
        return {}

    found: dict[str, Any] = {}
    byline = _clean(_decode_iptc(info.get(_IPTC_BYLINE)))
    if byline:
        found["creator"] = byline

    source = _clean(_decode_iptc(info.get(_IPTC_SOURCE)))
    if source:
        found["publisher"] = source

    credit = _clean(_decode_iptc(info.get(_IPTC_CREDIT)))
    if credit:
        found["credit_line"] = credit

    return found


def _decode_iptc(value: Any) -> Optional[str]:
    """IPTC values are bytes, and repeat-able fields come back as a list of them."""
    if isinstance(value, (list, tuple)):
        value = value[0] if value else None
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="replace")
    return value if isinstance(value, str) else None


# ─── documents ───────────────────────────────────────────────────────────────


def _from_document(path: Path, original_name: str) -> dict[str, Any]:
    extension = extension_of(original_name) or extension_of(path.name)

    if extension == ".pdf":
        return _from_pdf(path)
    if extension == ".docx":
        return _from_docx(path)
    if extension == ".pptx":
        return _from_pptx(path)
    if extension == ".xlsx":
        return _from_xlsx(path)
    # .txt/.md/.csv carry no metadata, and the legacy formats extract_text already
    # refuses by name are not readable here either.
    return {}


def _from_pdf(path: Path) -> dict[str, Any]:
    import pypdfium2

    document = pypdfium2.PdfDocument(path)
    try:
        return _from_core_properties(
            author=document.get_metadata_value("Author"),
            title=document.get_metadata_value("Title"),
            created=document.get_metadata_value("CreationDate"),
        )
    finally:
        document.close()


def _from_docx(path: Path) -> dict[str, Any]:
    import docx

    properties = docx.Document(str(path)).core_properties
    return _from_core_properties(
        author=properties.author,
        title=properties.title,
        created=properties.created,
        publisher=properties.company if hasattr(properties, "company") else None,
    )


def _from_pptx(path: Path) -> dict[str, Any]:
    from pptx import Presentation

    properties = Presentation(str(path)).core_properties
    return _from_core_properties(
        author=properties.author,
        title=properties.title,
        created=properties.created,
    )


def _from_xlsx(path: Path) -> dict[str, Any]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        properties = workbook.properties
        return _from_core_properties(
            author=properties.creator,
            title=properties.title,
            created=properties.created,
        )
    finally:
        workbook.close()


def _from_core_properties(
    *,
    author: Any = None,
    title: Any = None,
    created: Any = None,
    publisher: Any = None,
) -> dict[str, Any]:
    """The shape every document format reduces to.

    Unlike a media container, a document's `title` *is* the work's own title, so it maps
    to `source_title`. See the note on `_TAG_SOURCES` for why media goes the other way.
    """
    found: dict[str, Any] = {}

    cleaned_author = _clean(author)
    # Office writes "python-docx"-style generator names into `author` when nobody set
    # one, and Word defaults it to the machine's registered owner. Neither is worth
    # rejecting heuristically — the user can clear it, and a blank field teaches them
    # nothing — but an empty-ish one is not worth writing either.
    if cleaned_author:
        found["creator"] = cleaned_author

    cleaned_title = _clean(title)
    if cleaned_title:
        found["source_title"] = cleaned_title

    cleaned_publisher = _clean(publisher)
    if cleaned_publisher:
        found["publisher"] = cleaned_publisher

    when = normalise_partial_date(created)
    if when:
        found["published_date"] = when

    return found


# ─── shared helpers ──────────────────────────────────────────────────────────

# Leading year, then optionally month and day, separated by - or : (EXIF uses colons).
_DATE_HEAD = re.compile(r"^\s*(\d{4})(?:[-:/](\d{1,2}))?(?:[-:/](\d{1,2}))?")

# PDF writes a separator-less date with a marker prefix: "D:20190315101112Z". Matched
# separately because the run-together digits are ambiguous under the pattern above —
# it would read the year and then stop, losing the month and day.
_DATE_COMPACT = re.compile(r"^(?:D:)?(\d{4})(\d{2})(\d{2})(?:\d{2})*[Zz+\-]?")


def normalise_partial_date(value: Any) -> Optional[str]:
    """Reduce whatever a file claims as a date to `YYYY`, `YYYY-MM` or `YYYY-MM-DD`.

    Containers are wildly inconsistent here: MP4 writes an ISO instant
    ("2019-03-15T10:00:00.000000Z"), ID3 often writes a bare year, EXIF uses colons
    ("2019:03:15 10:00:00"), and Office hands back a real datetime object. All of them
    reduce to the partial-date string `Asset.published_date` stores.

    Precision is never invented: a file that says only "2019" yields "2019", not
    "2019-01-01". That is the entire reason the column is a string — see the comment on
    the model field.
    """
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")

    raw = str(value).strip()

    compact = _DATE_COMPACT.match(raw)
    if compact:
        year, month, day = compact.groups()
        return _assemble_date(year, month, day)

    match = _DATE_HEAD.match(raw.removeprefix("D:"))
    if not match:
        return None

    year, month, day = match.groups()
    return _assemble_date(year, month, day)


def _assemble_date(
    year: str, month: Optional[str], day: Optional[str]
) -> Optional[str]:
    """Build the widest valid partial date these parts support.

    Degrades rather than rejects: a file claiming month 19 still knows its year, and
    keeping the year is better than discarding the whole field over one bad component.
    """
    if not 1000 <= int(year) <= 9999:
        return None
    if month is None:
        return year
    if not 1 <= int(month) <= 12:
        return year
    if day is None:
        return f"{year}-{int(month):02d}"
    if not 1 <= int(day) <= 31:
        return f"{year}-{int(month):02d}"
    return f"{year}-{int(month):02d}-{int(day):02d}"


def _looks_like_url(value: str) -> bool:
    return value.lower().startswith(("http://", "https://"))


def _clean(value: Any) -> Optional[str]:
    """Trim, and treat whitespace-only and Pillow's trailing NULs as absent."""
    if value is None:
        return None
    if isinstance(value, bytes):
        value = value.decode("utf-8", errors="replace")
    if not isinstance(value, str):
        return None
    cleaned = value.replace("\x00", "").strip()
    return cleaned or None
